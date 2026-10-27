import argparse
import os
from pathlib import Path
import re
import sys
from typing import Any, List, Optional, Set, Tuple

# ตั้งค่า stdout รองรับ utf-8 บน Windows
if sys.platform.startswith("win"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

try:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
except ImportError:
    print("ไม่พบ openpyxl: กรุณารัน 'pip install -r requirements.txt'")
    if getattr(sys, "frozen", False):
        input("กด Enter เพื่อออก...")
    sys.exit(1)

YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def normalize_id(val: Any) -> Optional[str]:
    if val is None:
        return None
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    val_str = str(val).strip()
    return val_str or None


def sanitize_filename(name: str) -> str:
    clean = Path(name).name
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", clean).strip() or "output"


def find_column(ws, keyword: str = "รายการ", max_rows: int = 15) -> Tuple[Optional[int], Optional[int]]:
    target = keyword.strip().lower()
    for row in range(1, min(max_rows, ws.max_row or max_rows) + 1):
        for col in range(1, (ws.max_column or 50) + 1):
            val = ws.cell(row=row, column=col).value
            if val and target in str(val).strip().lower():
                return row, col
    return None, None


def delete_rows_batched(ws, rows_to_delete: List[int]):
    if not rows_to_delete:
        return
    rows = sorted(rows_to_delete)
    batches = []
    start = rows[0]
    count = 1
    for r in rows[1:]:
        if r == start + count:
            count += 1
        else:
            batches.append((start, count))
            start = r
            count = 1
    batches.append((start, count))

    for start_row, num in reversed(batches):
        ws.delete_rows(start_row, num)


class ExcelProcessor:
    def __init__(
        self,
        base_dir: Optional[Path] = None,
        sold_dir: Optional[Path] = None,
        all_data_dir: Optional[Path] = None,
        output_dir: Optional[Path] = None,
        column_name: str = "รายการ",
    ):
        if base_dir:
            self.base_dir = Path(base_dir).resolve()
        elif getattr(sys, "frozen", False):
            self.base_dir = Path(sys.executable).parent.resolve()
        else:
            self.base_dir = Path(__file__).parent.resolve()

        self.sold_dir = Path(sold_dir).resolve() if sold_dir else self.base_dir / "input_sold"
        self.all_data_dir = Path(all_data_dir).resolve() if all_data_dir else self.base_dir / "input_all_data"
        self.output_dir = Path(output_dir).resolve() if output_dir else self.base_dir / "output"
        self.column_name = column_name

    def read_sold_ids(self, path: Path) -> Set[str]:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            header_row, col_idx = find_column(ws, self.column_name)
            if not col_idx:
                raise ValueError(f"ไม่พบคอลัมน์ '{self.column_name}' ใน {path.name}")

            ids = set()
            for row in ws.iter_rows(min_row=header_row + 1, min_col=col_idx, max_col=col_idx, values_only=True):
                norm = normalize_id(row[0] if row else None)
                if norm:
                    ids.add(norm)
            return ids
        finally:
            wb.close()

    def process_file(self, sold_path: Path, master_path: Path) -> bool:
        print(f"\nประมวลผล: {sold_path.name}")
        try:
            ids_to_find = self.read_sold_ids(sold_path)
            if not ids_to_find:
                print("  - ข้าม: ไม่พบรายการในไฟล์")
                return False

            wb = load_workbook(master_path)
            try:
                ws = wb.active
                header_row, col_idx = find_column(ws, self.column_name)
                if not col_idx:
                    print(f"  - ข้าม: ไม่พบคอลัมน์ '{self.column_name}' ในฐานข้อมูลหลัก")
                    return False

                rows_to_delete = []
                matched_count = 0
                matched_ids = set()
                start_row = header_row + 1
                max_r = ws.max_row or start_row

                for r in range(start_row, max_r + 1):
                    val = normalize_id(ws.cell(row=r, column=col_idx).value)
                    if val in ids_to_find:
                        matched_count += 1
                        matched_ids.add(val)
                        for c in range(1, (ws.max_column or 50) + 1):
                            ws.cell(row=r, column=c).fill = YELLOW_FILL
                    else:
                        rows_to_delete.append(r)

                missing = ids_to_find - matched_ids
                if missing:
                    sample = list(missing)[:5]
                    print(f"  - แจ้งเตือน: พบ {len(missing)} รายการขายที่ไม่มีในฐานข้อมูลหลัก (เช่น {', '.join(sample)})")

                delete_rows_batched(ws, rows_to_delete)

                out_name = f"{sanitize_filename(sold_path.stem)}_processed.xlsx"
                out_path = self.output_dir / out_name
                wb.save(out_path)
                print(f"  - เรียบร้อย: ตรงกัน {matched_count} แถว, ลบ {len(rows_to_delete)} แถว -> {out_name}")
                return True
            finally:
                wb.close()

        except PermissionError:
            print(f"  - บันทึกไม่สำเร็จ: กรุณาปิดไฟล์ Excel ที่เปิดค้างไว้ใน output/ ก่อน")
            return False
        except Exception as e:
            print(f"  - ผิดพลาด: {e}")
            return False

    def run(self):
        self.sold_dir.mkdir(parents=True, exist_ok=True)
        self.all_data_dir.mkdir(parents=True, exist_ok=True)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        master_files = [f for f in self.all_data_dir.glob("*.xlsx") if not f.name.startswith("~$")]
        if not master_files:
            print("ไม่พบไฟล์ใน input_all_data/")
            return

        sold_files = [f for f in self.sold_dir.glob("*.xlsx") if not f.name.startswith("~$")]
        if not sold_files:
            print("ไม่พบไฟล์ใน input_sold/")
            return

        master_path = master_files[0]
        print(f"ฐานข้อมูลหลัก: {master_path.name}")
        print(f"ไฟล์ที่ต้องทำ: {len(sold_files)} ไฟล์")

        success = sum(1 for f in sold_files if self.process_file(f, master_path))

        print(f"\nเสร็จสิ้น: สำเร็จ {success}/{len(sold_files)} ไฟล์")
        print(f"ผลลัพธ์อยู่ที่: {self.output_dir}\n")


def parse_args():
    parser = argparse.ArgumentParser(description="SalesMatchPro - จับคู่และคัดกรองข้อมูล Excel")
    parser.add_argument("--sold-dir", type=str, help="โฟลเดอร์ไฟล์ขาย")
    parser.add_argument("--all-data-dir", type=str, help="โฟลเดอร์ฐานข้อมูลหลัก")
    parser.add_argument("--output-dir", type=str, help="โฟลเดอร์ผลลัพธ์")
    parser.add_argument("--column", type=str, default="รายการ", help="ชื่อคอลัมน์ที่ต้องการค้นหา (ค่าเริ่มต้น: รายการ)")
    parser.add_argument("--no-pause", action="store_true", help="ไม่ต้องรอกด Enter เมื่อทำงานเสร็จ")
    return parser.parse_args()


def main():
    args = parse_args()
    try:
        processor = ExcelProcessor(
            sold_dir=Path(args.sold_dir) if args.sold_dir else None,
            all_data_dir=Path(args.all_data_dir) if args.all_data_dir else None,
            output_dir=Path(args.output_dir) if args.output_dir else None,
            column_name=args.column,
        )
        processor.run()
    except KeyboardInterrupt:
        print("\nยกเลิกการทำงาน")
    except Exception as e:
        print(f"\nเกิดข้อผิดพลาด: {e}")
    finally:
        if not args.no_pause:
            try:
                input("กด Enter เพื่อปิดโปรแกรม...")
            except (EOFError, KeyboardInterrupt):
                pass


if __name__ == "__main__":
    main()
